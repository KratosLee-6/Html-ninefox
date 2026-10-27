"""PPTX 导入审核台 API 门禁（v0.8.3）。

镜像 test_docx_intake_api 的纪律：真实 HTTP → 解析 → 提案 → 采纳 → 真实
生成用例 → 磁盘产物，断言产物里是这份 PPT 自己的正文与讲稿，而不是演示
文案；解析边界（几何不还原）必须随响应到达客户端。
"""

from __future__ import annotations

import base64
import io
from pathlib import Path

import pytest
from pptx import Presentation
from test_server_project_api import request as api_request

from tests.conftest import WorkbenchServer

PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJ"
    "AAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==")


def build_pptx(*, title: str = "增长团队 Q4 汇报") -> bytes:
    """两页：封面＋要点页（带讲稿备注）。python-pptx 现写真实 OOXML。"""
    prs = Presentation()
    prs.core_properties.title = title

    cover = prs.slides.add_slide(prs.slide_layouts[0])
    cover.shapes.title.text = title
    cover.placeholders[1].text = "面向管理层的季度汇报"

    content = prs.slides.add_slide(prs.slide_layouts[1])
    content.shapes.title.text = "三步走"
    body = content.placeholders[1].text_frame
    body.text = "第一步：拉新，渠道投放到位"
    body.add_paragraph().text = "第二步：激活，首周留存拉起"
    content.notes_slide.notes_text_frame.text = "讲稿：强调预算与时间线的对应关系。"

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def upload_payload(data: bytes | None = None, **extra) -> dict:
    body = {"name": "deck.pptx"}
    if data is not None:
        body["pptx_base64"] = base64.b64encode(data).decode("ascii")
    body.update(extra)
    return body


@pytest.fixture()
def workbench(tmp_path: Path):
    with WorkbenchServer(tmp_path / "output") as server:
        yield server


def test_upload_returns_a_reviewable_deck_proposal(workbench):
    payload, _ = api_request(workbench.base_url, "/api/intake/pptx", "POST",
                             upload_payload(build_pptx()))

    assert payload["source"]["kind"] == "pptx"
    assert payload["source"]["title"] == "增长团队 Q4 汇报"
    assert payload["counts"]["heading"] == 2
    assert payload["counts"]["list"] == 2
    assert payload["counts"]["paragraph"] == 2  # 封面副标题 + 讲稿
    proposal = payload["proposal"]
    assert proposal["intent"] == "deck", "pptx 来源的默认意图必须是 deck"
    assert "演示文稿" in proposal["intent_reason"]
    assert proposal["style"]["preset_id"]
    assert proposal["image_plan"], "提案必须带图片位"
    assert any("几何" in note for note in payload["limitations"]), \
        "解析边界必须随提案一起返回"
    assert payload["evidence_complete"] is True


@pytest.mark.parametrize("data,code", [
    (None, "pptx_upload_required"),
    (b"not a pptx at all", "pptx_unreadable"),
])
def test_upload_errors_name_the_reason(workbench, data, code):
    payload, _ = api_request(workbench.base_url, "/api/intake/pptx", "POST",
                             upload_payload(data), expected=400)

    assert payload["error"]["code"] == code


def test_upload_rejects_broken_base64(workbench):
    payload, _ = api_request(workbench.base_url, "/api/intake/pptx", "POST",
                             {"pptx_base64": "%%%not-base64%%%"}, expected=400)

    assert payload["error"]["code"] == "pptx_upload_invalid"


def test_adopt_renders_the_deck_prose_into_a_real_project(workbench):
    payload, _ = api_request(workbench.base_url, "/api/intake/pptx/adopt", "POST",
                             upload_payload(build_pptx()))

    assert payload["ok"] is True
    assert payload["pptx"]["sections"] >= 2
    html = (Path(payload["project"]) / "output.html").read_text(encoding="utf-8")

    assert "第一步：拉新，渠道投放到位" in html, "幻灯片要点要进产物"
    assert "讲稿：强调预算与时间线的对应关系。" in html, "讲稿备注也要进产物"
    assert "Your Product" not in html, "品牌为空时不能印占位符"
    assert any("几何" in item for item in payload["limitations"])


def test_review_desk_can_change_intent_and_preset_before_adopting(workbench):
    payload, _ = api_request(
        workbench.base_url, "/api/intake/pptx/adopt", "POST",
        upload_payload(build_pptx(),
                       proposal={"intent": "doc", "style": {"preset_id": "doc-clean"}}))

    assert payload["intent"] == "doc"
    assert payload["preset_id"] == "doc-clean"
    html = (Path(payload["project"]) / "output.html").read_text(encoding="utf-8")
    assert "第一步：拉新，渠道投放到位" in html


def test_adopt_refuses_an_intent_the_desk_cannot_render(workbench):
    payload, _ = api_request(
        workbench.base_url, "/api/intake/pptx/adopt", "POST",
        upload_payload(build_pptx(), proposal={"intent": "hologram"}), expected=400)

    assert payload["error"]["code"] == "pptx_proposal_invalid"


@pytest.mark.parametrize("intent", ["landing", "dashboard", "deck", "poster", "archdoc", "doc"])
def test_every_intent_survives_the_real_generation_chain(workbench, intent):
    """真链：HTTP → 解析 → 提案 → 生成用例 → 磁盘产物，六个意图逐个走一遍。"""
    payload, _ = api_request(
        workbench.base_url, "/api/intake/pptx/adopt", "POST",
        upload_payload(build_pptx(), proposal={"intent": intent}))

    assert payload["intent"] == intent
    html = (Path(payload["project"]) / "output.html").read_text(encoding="utf-8")
    assert "第一步：拉新，渠道投放到位" in html, f"{intent} 丢了幻灯片正文"
    assert payload["verification"]["ok"] is not False


def test_adopt_with_slot_image_attaches_real_asset(workbench):
    """图片位挂真实资产（v0.8.4 收口）：附图经检查入库并进入生成素材。"""
    import io as _io
    from PIL import Image
    buffer = _io.BytesIO()
    Image.new("RGB", (64, 32), (36, 62, 143)).save(buffer, format="PNG")
    image_b64 = base64.b64encode(buffer.getvalue()).decode()

    payload, _ = api_request(
        workbench.base_url, "/api/intake/pptx/adopt", "POST",
        upload_payload(build_pptx(), slot_images={"cover": image_b64}))

    assert payload["ok"] is True
    assert len(payload["slot_attachments"]) == 1
    attachment = payload["slot_attachments"][0]
    assert attachment["slot"] == "cover"
    assert attachment["width"] == 64 and attachment["height"] == 32
    assert attachment["dominant_color"].startswith("#")
    assert attachment["input_id"], "附图必须入库为输入素材参与生成"

    plan_cover = next(item for item in payload["proposal"]["image_plan"]
                      if item["slot"] == "cover")
    assert plan_cover["asset"]["dominant_color"] == attachment["dominant_color"], \
        "检查摘要必须回写到提案的对应图片位"


def test_adopt_rejects_a_broken_slot_image(workbench):
    payload, _ = api_request(
        workbench.base_url, "/api/intake/pptx/adopt", "POST",
        upload_payload(build_pptx(), slot_images={"cover": base64.b64encode(b"not an image").decode()}),
        expected=400)

    assert payload["error"]["code"] == "image_asset_invalid"
