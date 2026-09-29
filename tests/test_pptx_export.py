"""V0.6-S13: PPTX file bridge — controlled mapping from deck artifacts."""

from __future__ import annotations

import importlib.util
from pathlib import Path

import pytest

from htmlninefox import pipeline

pytest.importorskip("pptx", reason="python-pptx optional extra")


def _make_deck(tmp_path: Path) -> Path:
    return pipeline.run_expert(
        "做一个 AI 产品发布会 PPT，像素花园风格", intent_override="deck",
        output=str(tmp_path), quiet_llm=True)["work"]


def test_deck_artifact_exports_to_editable_pptx(tmp_path: Path) -> None:
    from htmlninefox import pptx_export

    work = _make_deck(tmp_path)
    report = pptx_export.export_deck_pptx(work)

    target = Path(report["file"])
    assert target.is_file() and target.suffix == ".pptx"
    assert report["slides"] >= 5, "生成的 deck 应有多个分页"
    assert isinstance(report["editable_elements"], int) and report["editable_elements"] > 0

    # 可编辑验证：用 python-pptx 重新打开，逐页存在文本框且标题非空
    from pptx import Presentation
    presentation = Presentation(str(target))
    assert len(presentation.slides) == report["slides"]
    first = presentation.slides[0]
    texts = [shape.text_frame.text for shape in first.shapes
             if shape.has_text_frame and shape.text_frame.text.strip()]
    assert texts, "首页应有可编辑文本"
    assert any(len(text) > 3 for text in texts)


def test_non_deck_artifact_is_rejected_with_clear_error(tmp_path: Path) -> None:
    from htmlninefox import pptx_export

    work = pipeline.run_expert(
        "做一个极简落地页", output=str(tmp_path), quiet_llm=True)["work"]
    with pytest.raises(RuntimeError, match="slide"):
        pptx_export.export_deck_pptx(work)


def test_parse_deck_slides_reports_flattened_visuals() -> None:
    from htmlninefox.pptx_export import parse_deck_slides
    html = (b"<html><body>"
            b'<section class="slide"><h1>One</h1><svg width="10"></svg></section>'
            b'<section class="slide"><h1>Two</h1></section></body></html>')
    slides = parse_deck_slides(html.decode("utf-8"))
    assert len(slides) == 2
    assert slides[0]["flattened"] == ["svg×1"]
    assert slides[1]["flattened"] == []


def test_deck_generation_stays_intent_faithful_with_or_without_jinja2(tmp_path: Path) -> None:
    """Regression (V0.6-S17): deck must render paginated slides in both environments.

    generate_expert used to alias deck → templates/landing.html, so any environment
    with the optional `templates` extra (jinja2) rendered a landing page with
    hero/features/pricing sections. The artifact then had no <section class="slide">,
    which silently broke the PPTX bridge, the slide edit API and the export center —
    i.e. the whole V0.6 G4 outcome. Deck now falls through to the native
    generators/deck.py renderer regardless of whether jinja2 is installed.
    """
    from htmlninefox.experts import generate_expert

    assert "deck" not in generate_expert._INTENT_TEMPLATE_ALIAS
    assert "deck" in generate_expert._INTENT_TEMPLATE_ALIAS_NO_FALLBACK

    work = _make_deck(tmp_path)
    html = (work / "output.html").read_text(encoding="utf-8")

    from htmlninefox.pptx_export import parse_deck_slides
    slides = parse_deck_slides(html)
    assert len(slides) >= 5, "deck 产物必须保持分页 slide 结构"
    assert slides[0]["title"], "首页标题可被提取"
    assert slides[0]["bullets"] or slides[0]["subtitle"], "首页应含可编辑正文"


# ---------------------------------------------------------------- S13U-a slide editing API


def test_slides_edit_api_creates_revision_and_conflict_protects(tmp_path: Path) -> None:
    import json
    import time
    import urllib.error
    import urllib.request

    from htmlninefox import revisions
    from htmlninefox.server import app as server_app
    from tests.conftest import WorkbenchServer

    work = pipeline.run_expert("做一个发布会 PPT", intent_override="deck",
                               output=str(tmp_path), quiet_llm=True)["work"]
    with WorkbenchServer(tmp_path) as server:
        base = server.base_url

        def call(path, method="GET", payload=None, expected=200):
            data = json.dumps(payload).encode() if payload is not None else None
            request = urllib.request.Request(base + path, data=data, method=method,
                                             headers={"Content-Type": "application/json"})
            try:
                with urllib.request.urlopen(request, timeout=15) as response:
                    assert response.status == expected
                    return json.loads(response.read().decode())
            except urllib.error.HTTPError as error:
                assert error.code == expected, error.read().decode()
                return json.loads(error.read().decode())

        slides_payload = call(f"/api/projects/{work.name}/slides")
        slides = slides_payload["slides"]
        current_revision = slides_payload["revision"]
        assert len(slides) >= 5
        assert slides[0]["texts"], "每个分页都应有可编辑文本节点"

        target = slides[1]["texts"][0]
        edits = [{"slide": target["slide"], "node": target["node"], "text": "编辑后的章节标题"}]

        conflict = call(f"/api/projects/{work.name}/slides", "PUT",
                        {"edits": edits, "expected_revision": current_revision + 99},
                        expected=409)
        assert conflict["error"]["code"] == "revision_conflict"

        empty = call(f"/api/projects/{work.name}/slides", "PUT",
                     {"edits": [], "expected_revision": current_revision}, expected=400)
        assert empty["error"]["code"] == "slides_edit_empty"

        result = call(f"/api/projects/{work.name}/slides", "PUT",
                      {"edits": edits, "expected_revision": current_revision})
        assert result["revision"] == current_revision + 1
        assert "编辑后的章节标题" in (work / "output.html").read_text(encoding="utf-8")
        # 生成新 Revision，历史保留
        history = [item["revision"] for item in revisions.history(work)]
        assert history == list(range(result["revision"] + 1))


# ---------------------------------------------------------------- S13U-b slides editor dialog


def test_slides_editor_dialog_edits_and_saves(tmp_path: Path) -> None:
    from playwright.sync_api import sync_playwright

    from htmlninefox.server import app as server_app
    from tests.conftest import WorkbenchServer

    work = pipeline.run_expert("做一个发布会 PPT", intent_override="deck",
                               output=str(tmp_path), quiet_llm=True)["work"]
    with WorkbenchServer(tmp_path) as server:
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch()
            page = browser.new_page(viewport={"width": 1440, "height": 900})
            errors: list[str] = []
            page.on("pageerror", lambda error: errors.append(str(error)))
            page.add_init_script("try{localStorage.clear()}catch(e){}")
            page.goto(server.base_url + "/")
            page.wait_for_function(
                "window.FoxInteraction && window.FoxSlides && nodes.length >= 5")
            node_id = page.evaluate(
                """project => {
                    const ws = activeWorkspace();
                    const node = addNode('output', ws.x + ws.w + 80, ws.y + 30, {
                        title:'发布会 PPT · 产物', project_name: project,
                        preview_url:'/output/' + project + '/output.html',
                        intent:'deck', preset_id:'fox-pixel-garden',
                        revision:0, feedback:[], workspaceId:ws.id,
                    });
                    select(node.id);
                    return node.id;
                }""", work.name)

            page.evaluate("nodeId => openSlideEditor(nodeId)", node_id)
            page.wait_for_selector("#slides-modal:not([hidden])")
            page.wait_for_function(
                "document.querySelectorAll('#slides-editor textarea').length >= 5")

            first_area = page.locator("#slides-editor textarea").first
            original = first_area.input_value()
            first_area.fill(original + "（已编辑）")
            page.locator("#slides-save").click()
            page.wait_for_function(
                "document.querySelector('.fox-toast[data-toast-type=success]')?.textContent.includes('幻灯片已更新')")
            page.screenshot(path=str(tmp_path / "slides-editor-save.png"))

            # 版本徽标与产物内容双确认
            assert "rev1" in page.locator("#ins-rev").inner_text()
            assert "（已编辑）" in (work / "output.html").read_text(encoding="utf-8")
            assert errors == []
            browser.close()
