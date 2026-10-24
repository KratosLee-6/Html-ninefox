"""pptx intake gates (v0.8.3).

The fixture builds a real .pptx with python-pptx itself — the same writer the
export side uses — so these tests exercise the same OOXML an upload takes:
no hand-built XML, no network, no binary fixture in the repo nobody can
inspect. python-pptx is a dev dependency; it deliberately does not skip on
ImportError (same discipline as cryptography in test_ssrf_tls_pinning).

The contract under test is shared (`intake_outline`), not PPTX-shaped: every
block carries a locator (slide + shape), every outline carries its
limitations, and sampled locators are followed back into the source file and
compared — evidence you cannot walk back is not evidence.
"""

from __future__ import annotations

import base64
import io
import zipfile

import pytest
from pptx import Presentation
from pptx.util import Inches

from htmlninefox.intake_outline import Block, IntakeOutline, IntakeOutlineError
from htmlninefox.pptx_import import GEOMETRY_LIMITATION, PptxIntakeError, parse_pptx

# 1×1 透明 PNG：add_picture 要吃一个它认识的图格式。
PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJ"
    "AAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==")


def build_pptx(tmp_path, *, core_title: str = ""):
    """三页真实幻灯片：封面、要点页（带讲稿）、表格＋图片页。"""
    prs = Presentation()
    if core_title:
        prs.core_properties.title = core_title

    cover = prs.slides.add_slide(prs.slide_layouts[0])
    cover.shapes.title.text = "产品发布计划"
    cover.placeholders[1].text = "面向增长团队的路演"

    content = prs.slides.add_slide(prs.slide_layouts[1])
    content.shapes.title.text = "三步走"
    body = content.placeholders[1].text_frame
    body.text = "第一步：拉新"
    item = body.add_paragraph()
    item.text = "渠道投放"
    item.level = 1
    body.add_paragraph().text = "第二步：激活"
    content.notes_slide.notes_text_frame.text = "强调渠道预算与时间线。"

    board = prs.slides.add_slide(prs.slide_layouts[5])
    board.shapes.title.text = "预算盘子"
    table = board.shapes.add_table(2, 2, Inches(1), Inches(2),
                                   Inches(8), Inches(1.5)).table
    for row_index, row in enumerate((("渠道", "预算"), ("搜索", "40%"))):
        for column, value in enumerate(row):
            table.cell(row_index, column).text = value
    board.shapes.add_picture(io.BytesIO(PNG), Inches(1), Inches(4),
                             width=Inches(2))

    path = tmp_path / "deck.pptx"
    prs.save(path)
    return path


def test_outline_carries_titles_bullets_tables_and_pictures(tmp_path):
    outline = parse_pptx(build_pptx(tmp_path))

    assert outline.source.kind == "pptx"
    assert outline.source.title == "产品发布计划", "核心属性为空时回退首页标题"
    assert outline.source.digest.startswith("sha256:")
    assert outline.source.origin == "user_upload" and outline.source.licence == "user_owned"

    kinds = [b.kind for b in outline.blocks]
    assert kinds.count("heading") == 3, kinds
    assert kinds.count("list") == 3
    assert kinds.count("paragraph") == 2  # 封面副标题 + 讲稿备注
    assert kinds.count("table") == 1
    assert kinds.count("image") == 1

    headings = [b.heading for b in outline.blocks if b.kind == "heading"]
    assert headings == ["产品发布计划", "三步走", "预算盘子"]

    lists = [b for b in outline.blocks if b.kind == "list"]
    assert [b.level for b in lists] == [1, 2, 1], "缩进级别 ilvl → 深度 1/2/1"

    table = next(b for b in outline.blocks if b.kind == "table")
    assert table.rows == (("渠道", "预算"), ("搜索", "40%"))


def test_core_properties_title_wins_over_the_cover(tmp_path):
    outline = parse_pptx(build_pptx(tmp_path, core_title="增长团队 Q4 汇报"))
    assert outline.source.title == "增长团队 Q4 汇报"


def test_every_block_carries_a_locator(tmp_path):
    """没有可回指位置的块不允许进审核台——每个块至少带 slide 号。"""
    outline = parse_pptx(build_pptx(tmp_path))

    assert outline.evidence_complete is True
    for block in outline.blocks:
        assert isinstance(block.evidence.get("slide"), int), f"{block.id} 没有 slide 定位"
        assert block.evidence["slide"] >= 1


def test_locators_walk_back_to_the_source(tmp_path):
    """抽查三条 locator，回到源文件按位置取出原文比对——evidence 必须走得回去。"""
    path = build_pptx(tmp_path)
    outline = parse_pptx(path)
    prs = Presentation(str(path))

    heading = next(b for b in outline.blocks if b.kind == "heading" and b.heading == "三步走")
    assert heading.evidence == {"slide": 2, "shape": 1}
    source_slide = prs.slides[heading.evidence["slide"] - 1]
    assert source_slide.shapes.title.text == heading.heading

    bullet = next(b for b in outline.blocks if b.kind == "list" and b.content == "渠道投放")
    assert bullet.evidence == {"slide": 2, "shape": 2, "para": 2}
    source_frame = prs.slides[bullet.evidence["slide"] - 1].placeholders[1].text_frame
    assert source_frame.paragraphs[bullet.evidence["para"] - 1].text == bullet.content

    table = next(b for b in outline.blocks if b.kind == "table")
    assert table.evidence == {"slide": 3, "shape": 2}
    source_table = list(prs.slides[table.evidence["slide"] - 1].shapes)[
        table.evidence["shape"] - 1].table
    assert source_table.cell(0, 0).text == table.rows[0][0]


def test_notes_become_a_paragraph_with_a_notes_locator(tmp_path):
    outline = parse_pptx(build_pptx(tmp_path))

    notes = next(b for b in outline.blocks if b.evidence.get("part") == "notes")
    assert notes.kind == "paragraph"
    assert notes.content == "强调渠道预算与时间线。"
    assert notes.evidence["slide"] == 2


def test_picture_becomes_an_asset_candidate(tmp_path):
    outline = parse_pptx(build_pptx(tmp_path))

    assert list(outline.assets) == ["pptx/slide-3-sh3.png"]
    assert outline.assets["pptx/slide-3-sh3.png"] == PNG
    image = next(b for b in outline.blocks if b.kind == "image")
    assert image.asset == "pptx/slide-3-sh3.png"


def test_machine_form_is_stable_for_external_agents(tmp_path):
    payload = parse_pptx(build_pptx(tmp_path)).as_dict()

    assert payload["source"]["kind"] == "pptx"
    assert payload["source"]["title"] == "产品发布计划"
    assert payload["adapter"] == {"name": "pptx", "version": 1, "offline": True,
                                  "requires": ["python-pptx"]}
    assert payload["counts"] == {"heading": 3, "paragraph": 2, "list": 3,
                                 "table": 1, "image": 1}
    assert payload["assets"] == [{"part": "pptx/slide-3-sh3.png", "bytes": len(PNG)}]
    assert payload["evidence_complete"] is True
    first = payload["blocks"][0]
    assert first["kind"] == "heading" and first["heading"] == "产品发布计划"
    assert first["level"] == 1 and first["evidence"] == {"slide": 1, "shape": 1}


def test_geometry_boundary_is_always_stated(tmp_path):
    """几何不还原是适配器边界，不是某次失败——它必须在每份结果里。"""
    assert parse_pptx(build_pptx(tmp_path)).limitations == (GEOMETRY_LIMITATION,)


def test_merged_cells_are_reported_not_hidden(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "合并示例"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2),
                                   Inches(8), Inches(1.5)).table
    for row_index, row in enumerate((("渠道", "预算"), ("搜索", "40%"))):
        for column, value in enumerate(row):
            table.cell(row_index, column).text = value
    table.cell(0, 0).merge(table.cell(0, 1))
    path = tmp_path / "merged.pptx"
    prs.save(path)

    outline = parse_pptx(path)
    assert any("合并单元格" in note for note in outline.limitations), outline.limitations
    table_block = next(b for b in outline.blocks if b.kind == "table")
    # python-pptx 的 merge 语义：起点格的文本是被合并格文本的拼接——
    # 解析器如实上报单元格现状，不猜原状（那是几何还原，属于不做的事）。
    assert "渠道" in table_block.rows[0][0] and "预算" in table_block.rows[0][0]


def test_charts_are_counted_not_decomposed(tmp_path):
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "趋势"
    data = CategoryChartData()
    data.categories = ("一月", "二月")
    data.add_series("拉新", (120, 180))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,
                           Inches(1), Inches(2), Inches(6), Inches(4), data)
    path = tmp_path / "chart.pptx"
    prs.save(path)

    outline = parse_pptx(path)
    assert [b.kind for b in outline.blocks] == ["heading"], "图表数据不得混进 blocks"
    assert any("图表" in note and "只计数" in note for note in outline.limitations)


def test_accepts_bytes_as_well_as_a_path(tmp_path):
    path = build_pptx(tmp_path)
    assert parse_pptx(path.read_bytes()).counts == parse_pptx(path).counts


def test_unreadable_upload_says_so_instead_of_half_parsing(tmp_path):
    with pytest.raises(PptxIntakeError) as excinfo:
        parse_pptx(b"not a pptx at all")
    assert "pptx" in str(excinfo.value)


def test_a_docx_uploaded_by_mistake_is_rejected(tmp_path):
    """.docx 也是 zip——好 zip 不等于 pptx，这里必须如实拒绝。"""
    path = tmp_path / "mistake.pptx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "<w/>")

    with pytest.raises(PptxIntakeError):
        parse_pptx(path)


def test_empty_deck_reports_no_blocks(tmp_path):
    path = tmp_path / "empty.pptx"
    Presentation().save(path)

    outline = parse_pptx(path)
    assert outline.blocks == ()
    assert outline.evidence_complete is True  # 空集为真：没有块就没有断言
    assert any("没有可读的文本内容" in note for note in outline.limitations)
    assert any("没有标题" in note for note in outline.limitations)


def test_contract_rejects_a_block_without_evidence(tmp_path):
    outline = parse_pptx(build_pptx(tmp_path))
    broken = IntakeOutline(
        source=outline.source,
        adapter=outline.adapter,
        blocks=(Block(id="b1", kind="paragraph", content="没有定位"),),
        limitations=(),
    )
    with pytest.raises(IntakeOutlineError) as excinfo:
        broken.validate()
    assert "可回指" in str(excinfo.value)
