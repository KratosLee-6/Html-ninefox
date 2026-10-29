"""v0.8.4 发版截图：审核台 PDF 提案 + 图片位挂真实图片（真服务器实拍）。"""
from __future__ import annotations

import base64
import io
import sys
import tempfile
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

from playwright.sync_api import sync_playwright  # noqa: E402
from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

from tests.conftest import WorkbenchServer  # noqa: E402
from tests.test_pdf_import import build_pdf  # noqa: E402

OUT = ROOT / "assets" / "screenshots" / "v0.8.4"


def build_deck() -> bytes:
    prs = Presentation()
    prs.core_properties.title = "全域增长 Q4 汇报"
    cover = prs.slides.add_slide(prs.slide_layouts[0])
    cover.shapes.title.text = "全域增长 Q4 汇报"
    cover.placeholders[1].text = "面向管理层的季度汇报"
    plan = prs.slides.add_slide(prs.slide_layouts[1])
    plan.shapes.title.text = "三步走"
    body = plan.placeholders[1].text_frame
    body.text = "第一步：拉新"
    plan.notes_slide.notes_text_frame.text = "讲稿：预算与时间线。"
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def build_pdf_scanned() -> bytes:
    """两页全空内容流——扫描件的忠实替身。"""
    def stream_object(index: int) -> tuple[int, bytes]:
        content = b"BT ET\n"
        return index, (f"<</Length {len(content)}>>stream\n".encode("latin-1")
                       + content + b"endstream\n")

    body = b"%PDF-1.4\n"
    body += b"1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
    body += b"2 0 obj<</Type/Pages/Kids[3 0 R 5 0 R]/Count 2>>endobj\n"
    body += b"3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]/Contents 4 0 R>>endobj\n"
    _, stream4 = stream_object(4)
    body += b"4 0 obj" + stream4 + b"endobj\n"
    body += b"5 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]/Contents 6 0 R>>endobj\n"
    _, stream6 = stream_object(6)
    body += b"6 0 obj" + stream6 + b"endobj\n"
    body += b"trailer<</Root 1 0 R/Size 7>>\n%%EOF"
    return body


def build_png() -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (640, 360), (36, 62, 143)).save(buffer, format="PNG")
    return buffer.getvalue()


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    tmp = Path(tempfile.mkdtemp())
    pdf = tmp / "接入说明.pdf"
    pdf.write_bytes(build_pdf())
    deck = tmp / "全域增长Q4汇报.pptx"
    deck.write_bytes(build_deck())
    cover_png = tmp / "cover.png"
    cover_png.write_bytes(build_png())

    with WorkbenchServer(Path(tempfile.mkdtemp())) as server:
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch()
            page = browser.new_page(viewport={"width": 1440, "height": 1100})
            page.goto(server.base_url + "/")
            page.wait_for_function("window.FoxDocx && window.FoxIntake")
            page.evaluate("window.FoxIntake.open()")

            # shot 1: PDF 提案
            page.set_input_files("#docx-file", str(pdf))
            page.click("#docx-analyze")
            page.wait_for_selector("#docx-proposal:not([hidden])", timeout=30000)
            page.wait_for_timeout(300)
            page.locator("#docx-desk").screenshot(
                path=OUT / "pdf-desk-proposal-1440.png")
            print("shot: pdf-desk-proposal-1440.png")

            # shot 2: PPT 上传 → cover 槽位附图 → 摘要 chip
            # （语言切换钮在提案渲染后才可见，放在首次分析之后）
            page.set_input_files("#docx-file", str(deck))
            page.click("#docx-analyze")
            page.wait_for_selector("#docx-proposal:not([hidden])", timeout=30000)
            page.click("#docx-slot-lang")  # 切 EN，槽位名与 plan id 对齐
            page.locator("#docx-slots .docx-slot").first.locator(
                ".docx-slot-attach input").set_input_files(str(cover_png))
            page.wait_for_selector("#docx-slots [data-part='asset']:not([hidden])",
                                   timeout=30000)
            page.wait_for_timeout(300)
            page.locator("#docx-desk").screenshot(
                path=OUT / "slot-attach-1440.png")
            print("shot: slot-attach-1440.png")
            browser.close()

    # ── 场景 B：扫描件边界实况（无文本层的 PDF，边界清单如实显示）──
    scanned = tmp / "扫描件.pdf"
    scanned.write_bytes(build_pdf_scanned())
    with WorkbenchServer(Path(tempfile.mkdtemp())) as server:
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch()
            page = browser.new_page(viewport={"width": 1440, "height": 1100})
            page.goto(server.base_url + "/")
            page.wait_for_function("window.FoxDocx && window.FoxIntake")
            page.evaluate("window.FoxIntake.open()")
            page.set_input_files("#docx-file", str(scanned))
            page.click("#docx-analyze")
            page.wait_for_selector("#docx-proposal:not([hidden])", timeout=30000)
            page.wait_for_timeout(300)
            page.locator("#docx-desk").screenshot(
                path=OUT / "pdf-scanned-boundary-1440.png")
            print("shot: pdf-scanned-boundary-1440.png")
            browser.close()

    # ── 场景 C：PDF 采纳后的产物页（真实生成落盘）──
    with WorkbenchServer(Path(tempfile.mkdtemp())) as server:
        def post(path, body):
            import urllib.request
            req = urllib.request.Request(
                server.base_url + path, data=__import__("json").dumps(body).encode(),
                headers={"Content-Type": "application/json"}, method="POST")
            with urllib.request.urlopen(req, timeout=120) as r:
                return __import__("json").load(r)

        gen = post("/api/generate", {
            "prompt": "接入说明：把内容交付流水线接进团队流程，三步接入与格式对照",
            "intent": "doc", "quiet_llm": True})
        adopt = post("/api/intake/pdf/adopt", {
            "name": "接入说明.pdf", "pdf_base64": base64.b64encode(build_pdf()).decode(),
            "proposal": {"intent": "doc"}})
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch()
            page = browser.new_page(viewport={"width": 1440, "height": 1000})
            page.goto(server.base_url + adopt["preview_url"])
            page.wait_for_timeout(1800)
            page.screenshot(path=OUT / "pdf-artifact-1440.png")
            print("shot: pdf-artifact-1440.png")
            browser.close()
    print("capture done")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
