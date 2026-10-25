"""Capture v0.8.3 evidence shots: the two views this release adds.

  1. pptx-desk-proposal-1440.png — 审核台吃进 .pptx：提案、边界清单、可改表单
  2. pptx-artifact-1440.png     — 采纳后的产物（幻灯片要点与讲稿进了页面）

Both are real product states against a real server: a .pptx built with
python-pptx is uploaded through the same file input a user clicks, the
proposal is what /api/intake/pptx returned, and the artifact is the project
the adopt endpoint really wrote to disk.

    python scripts/capture_v083_shots.py
"""
from __future__ import annotations

import io
import sys
import tempfile
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from playwright.sync_api import sync_playwright  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

from tests.conftest import WorkbenchServer  # noqa: E402  (真服务器 fixture)

OUT = ROOT / "assets" / "screenshots" / "v0.8.3"


def build_deck() -> bytes:
    prs = Presentation()
    prs.core_properties.title = "全域增长 Q4 汇报"

    cover = prs.slides.add_slide(prs.slide_layouts[0])
    cover.shapes.title.text = "全域增长 Q4 汇报"
    cover.placeholders[1].text = "增长团队 · 面向管理层的季度汇报"

    plan = prs.slides.add_slide(prs.slide_layouts[1])
    plan.shapes.title.text = "三步走"
    body = plan.placeholders[1].text_frame
    body.text = "第一步：拉新，渠道投放到位"
    body.add_paragraph().text = "第二步：激活，首周留存拉起"
    body.add_paragraph().text = "第三步：变现，会员转化提效"
    plan.notes_slide.notes_text_frame.text = (
        "讲稿：三步的时间线与预算对应关系要讲透，激活是本季重心。")

    numbers = prs.slides.add_slide(prs.slide_layouts[5])
    numbers.shapes.title.text = "预算盘子"
    table = numbers.shapes.add_table(3, 2, Inches(1), Inches(2),
                                     Inches(8), Inches(2)).table
    for row_index, row in enumerate((("渠道", "预算"), ("搜索", "40%"),
                                     ("信息流", "35%"))):
        for column, value in enumerate(row):
            table.cell(row_index, column).text = value

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    deck = Path(tempfile.mkdtemp()) / "全域增长Q4汇报.pptx"
    deck.write_bytes(build_deck())

    with WorkbenchServer(Path(tempfile.mkdtemp())) as server:
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch()
            page = browser.new_page(viewport={"width": 1440, "height": 1100})
            page.goto(server.base_url + "/")
            page.wait_for_function("window.FoxDocx && window.FoxIntake")
            page.evaluate("window.FoxIntake.open()")

            page.set_input_files("#docx-file", str(deck))
            page.click("#docx-analyze")
            page.wait_for_selector("#docx-proposal:not([hidden])", timeout=20000)
            page.wait_for_timeout(400)
            page.locator("#docx-desk").screenshot(
                path=OUT / "pptx-desk-proposal-1440.png")
            print(f"shot 1: {OUT / 'pptx-desk-proposal-1440.png'}")

            page.click("#docx-adopt")
            page.wait_for_selector("#docx-result:not([hidden])", timeout=120000)
            href = page.get_attribute("#docx-result", "href")
            page.goto(server.base_url + href)
            page.wait_for_timeout(1200)
            page.screenshot(path=OUT / "pptx-artifact-1440.png", full_page=True)
            print(f"shot 2: {OUT / 'pptx-artifact-1440.png'}")
            browser.close()
    print("capture done")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
