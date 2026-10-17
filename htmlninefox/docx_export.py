"""DOCX semantic bridge (v0.8.0): controlled mapping from document artifacts.

Maps generated doc/archdoc HTML into a standards-compliant .docx via
python-docx, and deck HTML into a handout outline. Semantic mapping v1
(宁小勿假): h1–h6 → Heading 1–6, paragraphs, ul/ol list items, tables,
embedded images (data: URI), blockquotes, hr separators, pre/code as
monospace paragraphs. Everything the mapping drops is counted in the
report and never silently lost: layout containers' stray text, external
images (offline promise — never fetched), flattened visuals (deck), and
the global note that CSS colors/fonts/layout do not carry into Word.

Intent routing lives in exporting._export_docx: doc/archdoc → document
mapping, deck → handout outline, landing/poster/dashboard → StoreError
guidance toward PDF/PNG (Word is the wrong form for visual-first pages).
"""

from __future__ import annotations

import base64
import html as html_mod
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from html.parser import HTMLParser

TITLE_PATTERN = re.compile(r"<title[^>]*>([\s\S]*?)</title>", re.I)

# 语义块标签 → 块类型。透明容器（div/section/main/article…）只做递归上下文。
_HEADING_TAGS = {f"h{i}": i for i in range(1, 7)}
_BLOCK_OPEN = {"p": "paragraph", "blockquote": "quote", "pre": "code"}
_SKIP_TAGS = {"script", "style", "svg", "canvas", "video", "iframe", "noscript",
              "nav", "header", "footer", "aside", "template", "button",
              "select", "input", "textarea", "label", "form"}
# chrome 元素不是「被降级的内容」——是界面杂音，单独计数以示透明。
_CHROME_TAGS = {"nav", "header", "footer", "aside", "button", "select",
                "input", "textarea", "label", "form"}
_FLATTEN_TAGS = ("svg", "canvas", "video", "iframe")
_DATA_URI = re.compile(r"^data:image/(png|jpeg|jpg|gif|bmp);base64,(.*)$", re.S)


def _clean(text: str) -> str:
    return " ".join(html_mod.unescape(text).split())


class _SemanticMapper(HTMLParser):
    """One pass: semantic blocks out, every drop counted."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.title = ""
        self.blocks: list[dict[str, Any]] = []
        self.stack: list[str] = []
        self.current: dict[str, Any] | None = None
        self.current_tag = ""
        self.list_kinds: list[str] = []
        self.table: list[list[str]] | None = None
        self.rows: list[list[str]] | None = None
        self.cell: str | None = None
        self.cell_depth = 0
        self.in_title = False
        self.skip_stack: list[str] = []
        self.flattened: dict[str, int] = {}
        self.chrome: dict[str, int] = {}
        self.external_images: list[str] = []
        self.bad_images: int = 0
        self.unmapped_chars = 0

    # -- helpers ----------------------------------------------------------
    def _open(self, kind: str, tag: str, **extra) -> None:
        self.current = {"kind": kind, "text": "", **extra}
        self.current_tag = tag

    def _close(self) -> None:
        if self.current is None:
            return
        block = self.current
        kind = block["kind"]
        if kind == "code":
            block["text"] = block["text"].strip("\n").rstrip()
        else:
            block["text"] = _clean(block["text"])
        if not block.get("text") and kind in ("paragraph", "quote", "code"):
            self.current = None  # 空段不入文档，也不计数
            return
        self.blocks.append(block)
        self.current = None

    def _skip_active(self) -> bool:
        return bool(self.skip_stack)

    # -- parser hooks -----------------------------------------------------
    def handle_starttag(self, tag, attrs):
        tag = tag.lower()
        attrd = {name: (value or "") for name, value in attrs}
        if self._skip_active():
            if tag in _SKIP_TAGS:
                self.skip_stack.append(tag)
            return
        if tag in _SKIP_TAGS:
            self.skip_stack.append(tag)
            bucket = self.chrome if tag in _CHROME_TAGS else self.flattened
            bucket[tag] = bucket.get(tag, 0) + 1
            return
        if tag == "title":
            self.in_title = True
            return
        if tag == "img":
            src = attrd.get("src", "")
            match = _DATA_URI.match(src.strip())
            if match:
                self.blocks.append({"kind": "image", "data": match.group(2),
                                    "mime": "image/" + ("jpg" if match.group(1) == "jpeg" else match.group(1)),
                                    "alt": _clean(attrd.get("alt", ""))})
            elif src.strip():
                self.external_images.append(src.strip()[:120])
            return
        if tag == "hr":
            self.blocks.append({"kind": "separator"})
            return
        if tag in ("ul", "ol"):
            self.list_kinds.append("number" if tag == "ol" else "bullet")
            self.stack.append(tag)
            return
        if tag == "li":
            kind = self.list_kinds[-1] if self.list_kinds else "bullet"
            self._open("list_item", tag, style=kind)
            self.stack.append(tag)
            return
        if tag == "table":
            self.rows = []
            self.stack.append(tag)
            return
        if tag == "tr" and self.rows is not None:
            self.rows.append([])
            self.stack.append(tag)
            return
        if tag in ("td", "th") and self.rows is not None:
            if self.cell is not None:
                self.cell_depth += 1  # 嵌套单元格：文本并入外层（如实、罕见）
            else:
                self.cell = ""
            self.stack.append(tag)
            return
        if tag in _HEADING_TAGS:
            self._open("heading", tag, level=_HEADING_TAGS[tag])
            self.stack.append(tag)
            return
        if tag in _BLOCK_OPEN:
            self._open(_BLOCK_OPEN[tag], tag)
            self.stack.append(tag)
            return
        self.stack.append(tag)  # 透明容器

    def handle_endtag(self, tag):
        tag = tag.lower()
        if self._skip_active():
            if tag in self.skip_stack:
                while self.skip_stack and self.skip_stack.pop() != tag:
                    pass
            return
        if tag == "title":
            self.in_title = False
            return
        if tag in ("td", "th") and self.cell is not None:
            # 单元格收尾：文本入行。嵌套单元格（cell_depth）的文本已在
            # 累积进外层 cell，这里只退深度。
            if self.cell_depth:
                self.cell_depth -= 1
            else:
                if self.rows and self.rows[-1] is not None:
                    self.rows[-1].append(_clean(self.cell))
                self.cell = None
            while self.stack:
                if self.stack.pop() == tag:
                    break
            return
        if tag == "li" and self.current is not None and self.current["kind"] == "list_item":
            self._close()
        elif tag == "table" and self.rows is not None:
            rows = [row for row in self.rows if any(cell.strip() for cell in row)]
            if rows:
                self.blocks.append({"kind": "table", "rows": rows})
            self.rows = None
            self.cell = None
            self.cell_depth = 0
        if not self.stack:
            return
        # 弹栈直到匹配（容错畸形标签）
        while self.stack:
            open_tag = self.stack.pop()
            if open_tag in ("ul", "ol"):
                self.list_kinds.pop() if self.list_kinds else None
            if open_tag == tag:
                if self.current is not None and self.current_tag == tag:
                    self._close()
                break

    def handle_data(self, data):
        if self.in_title:
            self.title += data
            return
        if self._skip_active():
            return
        if self.cell is not None and self.rows is not None and self.rows:
            self.cell += data
            return
        if self.current is not None:
            self.current["text"] += data
            return
        if data.strip():
            self.unmapped_chars += len(data.strip())

    def close(self):
        super().close()
        if self.cell is not None and self.rows is not None:
            self.rows[-1].append(_clean(self.cell))
        self._close()


def parse_document_blocks(html_text: str) -> dict[str, Any]:
    """Semantic blocks + an honest account of everything the mapping drops."""
    mapper = _SemanticMapper()
    mapper.feed(html_text)
    mapper.close()
    title = _clean(mapper.title)[:200]
    if not title:
        first_heading = next((b["text"] for b in mapper.blocks
                              if b["kind"] == "heading"), "")
        title = first_heading[:200]
    return {
        "title": title,
        "blocks": mapper.blocks,
        "flattened": mapper.flattened,
        "chrome": mapper.chrome,
        "external_images": mapper.external_images,
        "bad_images": mapper.bad_images,
        "unmapped_chars": mapper.unmapped_chars,
    }


def _counts_of(blocks: list[dict[str, Any]]) -> dict[str, int]:
    return {
        "headings": sum(1 for b in blocks if b["kind"] == "heading"),
        "paragraphs": sum(1 for b in blocks if b["kind"] == "paragraph"),
        "list_items": sum(1 for b in blocks if b["kind"] == "list_item"),
        "tables": sum(1 for b in blocks if b["kind"] == "table"),
        "images_embedded": sum(1 for b in blocks if b["kind"] == "image"),
        "quotes": sum(1 for b in blocks if b["kind"] == "quote"),
        "separators": sum(1 for b in blocks if b["kind"] == "separator"),
        "code_blocks": sum(1 for b in blocks if b["kind"] == "code"),
    }


def _add_separator(doc) -> None:
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    paragraph = doc.add_paragraph()
    ppr = paragraph._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "6"),
                       ("w:space", "1"), ("w:color", "999999")):
        bottom.set(qn(key), value)
    border.append(bottom)
    ppr.append(border)


def _add_code_paragraph(doc, text: str) -> None:
    from docx.shared import Pt
    paragraph = doc.add_paragraph()
    lines = text.splitlines() or [""]
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        if index < len(lines) - 1:
            run.add_break()


def _add_image(doc, block: dict[str, Any]) -> bool:
    import io
    from docx.shared import Inches
    try:
        raw = base64.b64decode(block["data"])
        doc.add_picture(io.BytesIO(raw), width=Inches(6))
    except Exception:  # noqa: BLE001 - 损坏的图片数据如实降级，不让整个导出失败
        return False
    if block.get("alt"):
        caption = doc.add_paragraph(block["alt"])
        caption.runs[0].italic = True
    return True


def _render_document(doc, blocks: list[dict[str, Any]]) -> dict[str, int]:
    embedded_failures = 0
    for block in blocks:
        kind = block["kind"]
        if kind == "heading":
            doc.add_heading(block["text"], level=min(block["level"], 9))
        elif kind == "paragraph":
            doc.add_paragraph(block["text"])
        elif kind == "list_item":
            doc.add_paragraph(block["text"],
                              style="List Number" if block["style"] == "number"
                              else "List Bullet")
        elif kind == "quote":
            doc.add_paragraph(block["text"], style="Intense Quote")
        elif kind == "code":
            _add_code_paragraph(doc, block["text"])
        elif kind == "separator":
            _add_separator(doc)
        elif kind == "table":
            rows = block["rows"]
            table = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            table.style = "Table Grid"
            for r_index, row in enumerate(rows):
                for c_index, cell in enumerate(row):
                    table.cell(r_index, c_index).text = cell
        elif kind == "image":
            if not _add_image(doc, block):
                embedded_failures += 1
    counts = _counts_of(blocks)
    counts["images_failed"] = embedded_failures
    return counts


def export_document_docx(project: str | Path, out_dir: str | Path | None = None) -> dict[str, Any]:
    """Full semantic mapping for doc/archdoc artifacts."""
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - guarded by optional extra
        raise RuntimeError("需要 python-docx：pip install 'htmlninefox[docx]'") from exc

    project = Path(project)
    html_path = project / "output.html"
    if not html_path.is_file():
        raise FileNotFoundError(f"项目缺少 output.html：{project}")
    parsed = parse_document_blocks(html_path.read_text(encoding="utf-8", errors="replace"))
    if not parsed["blocks"]:
        raise RuntimeError("该产物没有可映射的语义内容，请使用 PDF / PNG 导出")

    out_dir = Path(out_dir) if out_dir else project / "exports" / (
        datetime.now().strftime("%Y%m%d-%H%M%S") + "-docx")
    out_dir.mkdir(parents=True, exist_ok=True)
    target = out_dir / f"{project.name}.docx"

    doc = Document()
    doc.core_properties.title = parsed["title"] or project.name
    counts = _render_document(doc, parsed["blocks"])
    doc.save(target)

    degraded: dict[str, Any] = {
        "external_images": parsed["external_images"],
        "unmapped_text_chars": parsed["unmapped_chars"],
        "flattened": dict(parsed["flattened"]),
        "chrome_skipped": dict(parsed["chrome"]),
        "images_failed": counts.pop("images_failed"),
    }
    report = {
        "file": str(target),
        "mode": "document",
        "title": parsed["title"],
        "mapped": counts,
        "degraded": degraded,
        "notes": [
            "CSS 配色/字体/布局不映射——Word 由 Normal/Heading 样式与内嵌图片承载内容。",
            "外链图片不抓取（离线承诺），仅 data: URI 内嵌图进入文档。",
            "链接降为纯文本；表格样式统一为 Table Grid。",
        ],
        "generated_at": datetime.now().isoformat(timespec="seconds"),
    }
    return report


def export_deck_handout_docx(project: str | Path, out_dir: str | Path | None = None) -> dict[str, Any]:
    """Handout outline for deck artifacts (reuses the PPTX slide parser)."""
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - guarded by optional extra
        raise RuntimeError("需要 python-docx：pip install 'htmlninefox[docx]'") from exc

    from . import pptx_export

    project = Path(project)
    html_path = project / "output.html"
    if not html_path.is_file():
        raise FileNotFoundError(f"项目缺少 output.html：{project}")
    html_text = html_path.read_text(encoding="utf-8", errors="replace")
    slides = pptx_export.parse_deck_slides(html_text)
    if not slides:
        raise RuntimeError("该产物不是分页 deck（未找到 slide 结构），请使用 PDF / PNG 导出")

    out_dir = Path(out_dir) if out_dir else project / "exports" / (
        datetime.now().strftime("%Y%m%d-%H%M%S") + "-docx")
    out_dir.mkdir(parents=True, exist_ok=True)
    target = out_dir / f"{project.name}-讲义.docx"

    doc = Document()
    title_match = TITLE_PATTERN.search(html_text)
    doc.core_properties.title = _clean(title_match.group(1)) if title_match else project.name
    doc.add_heading(doc.core_properties.title or project.name, level=1)
    for slide in slides:
        doc.add_heading(f"第 {slide['index'] + 1} 页 · {slide['title']}", level=2)
        if slide["subtitle"]:
            paragraph = doc.add_paragraph(slide["subtitle"])
            paragraph.runs[0].italic = True
        for bullet in slide["bullets"]:
            doc.add_paragraph(bullet, style="List Bullet")
    doc.save(target)

    flattened_all = [f"slide {s['index'] + 1}: {item}"
                     for s in slides for item in s["flattened"]]
    report = {
        "file": str(target),
        "mode": "deck-handout",
        "slides": len(slides),
        "mapped": {"headings": 1 + len(slides), "list_items": sum(len(s["bullets"]) for s in slides)},
        "degraded": {"flattened": flattened_all},
        "notes": ["讲义式大纲：每页一个二级标题 ＋ 要点列表；页面视觉（背景/图形）不在 Word 中重现，见扁平化清单。"],
        "generated_at": datetime.now().isoformat(timespec="seconds"),
    }
    return report
