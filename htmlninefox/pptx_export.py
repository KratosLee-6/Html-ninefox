"""PPTX file bridge (V0.6-S13): controlled mapping from deck artifacts.

Maps our generated deck HTML (section.slide pages) into a standards-
compliant .pptx via python-pptx: per slide a title text frame, body
paragraphs, and a solid background color. Complex visuals (svg, canvas,
video) are counted as flattened and reported, never silently dropped.
"""

from __future__ import annotations

import html as html_mod
import re
from datetime import datetime
from pathlib import Path
from typing import Any

SLIDE_PATTERN = re.compile(r'<section[^>]*class="[^"]*slide[^"]*"[^>]*>([\s\S]*?)</section>', re.I)
HEADING_PATTERN = re.compile(r"<h([12])[^>]*>([\s\S]*?)</h\1>", re.I)
PARAGRAPH_PATTERN = re.compile(r"<(p|li)[^>]*>([\s\S]*?)</\1>", re.I)
BG_PATTERN = re.compile(r"background(?:-color)?\s*:\s*(#[0-9a-fA-F]{3,8})")
FLATTEN_TAGS = ("svg", "canvas", "video", "iframe")


def _strip_tags(fragment: str) -> str:
    text = re.sub(r"<[^>]+>", " ", fragment)
    return " ".join(html_mod.unescape(text).split())


def _bg_of(slide_html: str, default: str) -> tuple[str, list[str]]:
    flattened: list[str] = []
    for tag in FLATTEN_TAGS:
        count = len(re.findall(rf"<{tag}[\s>]", slide_html, re.I))
        if count:
            flattened.append(f"{tag}×{count}")
    match = BG_PATTERN.search(slide_html)
    color = match.group(1) if match else default
    if len(color) == 4:
        color = "#" + "".join(ch * 2 for ch in color[1:])
    return color.upper(), flattened


def _luminance(hex_color: str) -> float:
    value = hex_color.lstrip("#")
    r, g, b = (int(value[i:i + 2], 16) / 255 for i in (0, 2, 4))
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def parse_deck_slides(html_text: str) -> list[dict[str, Any]]:
    """Extract controlled-mapping slide data from a generated deck artifact."""
    slides: list[dict[str, Any]] = []
    matches = SLIDE_PATTERN.findall(html_text)
    default_bg = "#173C8F"
    bg_match = re.search(r"body\s*{[^}]*background[^:]*:\s*(#[0-9a-fA-F]{3,6})", html_text, re.I)
    if bg_match:
        default_bg = bg_match.group(1)
    for index, slide_html in enumerate(matches):
        background, flattened = _bg_of(slide_html, default_bg)
        headings = [(int(level), _strip_tags(text)) for level, text in HEADING_PATTERN.findall(slide_html)]
        headings = [(level, text) for level, text in headings if text]
        bullets = [_strip_tags(text) for _, text in PARAGRAPH_PATTERN.findall(slide_html)]
        bullets = [text for text in bullets if text][:14]
        title = headings[0][1] if headings else f"第 {index + 1} 页"
        subtitle = headings[1][1] if len(headings) > 1 else ""
        slides.append({
            "index": index,
            "title": title[:120],
            "subtitle": subtitle[:120],
            "bullets": bullets,
            "background": background,
            "dark_background": _luminance(background) < 0.45,
            "flattened": flattened,
        })
    return slides


def export_deck_pptx(project: str | Path, out_dir: str | Path | None = None) -> dict[str, Any]:
    """Render the project's deck artifact into an editable .pptx file."""
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.util import Inches, Pt
    except ImportError as exc:  # pragma: no cover - guarded by optional extra
        raise RuntimeError("需要 python-pptx：pip install 'htmlninefox[pptx]'") from exc

    project = Path(project)
    html_path = project / "output.html"
    if not html_path.is_file():
        raise FileNotFoundError(f"项目缺少 output.html：{project}")
    html_text = html_path.read_text(encoding="utf-8", errors="replace")
    slides = parse_deck_slides(html_text)
    if not slides:
        raise RuntimeError("该产物不是分页 deck（未找到 slide 结构），请使用 PDF / PNG 导出")

    out_dir = Path(out_dir) if out_dir else project / "exports" / (
        datetime.now().strftime("%Y%m%d-%H%M%S") + "-pptx")
    out_dir.mkdir(parents=True, exist_ok=True)
    target = out_dir / f"{project.name}.pptx"

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    total = len(slides)
    flattened_all: list[str] = []

    for slide_data in slides:
        slide = prs.slides.add_slide(blank)
        bg_hex = slide_data["background"].lstrip("#")
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(bg_hex)
        text_color = "F4F0E7" if slide_data["dark_background"] else "17283D"
        accent = "49B894" if not slide_data["dark_background"] else "49B894"

        title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.6), Inches(12), Inches(1.3))
        frame = title_box.text_frame
        frame.word_wrap = True
        run = frame.paragraphs[0].add_run()
        run.text = slide_data["title"]
        run.font.size = Pt(40)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(text_color)

        body_lines: list[str] = []
        if slide_data["subtitle"]:
            body_lines.append(slide_data["subtitle"])
        body_lines.extend(slide_data["bullets"])
        if body_lines:
            body_box = slide.shapes.add_textbox(Inches(0.9), Inches(2.2), Inches(11.5), Inches(4.6))
            body_frame = body_box.text_frame
            body_frame.word_wrap = True
            for line_index, line in enumerate(body_lines):
                paragraph = body_frame.paragraphs[0] if line_index == 0 else body_frame.add_paragraph()
                run = paragraph.add_run()
                run.text = "• " + line if line_index else line
                run.font.size = Pt(20)
                run.font.color.rgb = RGBColor.from_string(text_color)

        footer = slide.shapes.add_textbox(Inches(11.6), Inches(7.0), Inches(1.6), Inches(0.4))
        run = footer.text_frame.paragraphs[0].add_run()
        run.text = f"{slide_data['index'] + 1} / {total}"
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor.from_string("8A93A6")
        flattened_all.extend(f"slide {slide_data['index'] + 1}: {item}" for item in slide_data["flattened"])

    prs.save(target)
    report = {
        "file": str(target),
        "slides": total,
        "editable_elements": sum(2 + len(s["bullets"]) for s in slides),
        "flattened_elements": flattened_all,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
    }
    return report


# ---------------------------------------------------------------- slide text editing (S13U-a)

from html.parser import HTMLParser as _HTMLParser


class _TextNodeMapper(_HTMLParser):
    """Map every visible text node to its exact char span, per slide."""

    SKIP = {"script", "style"}

    def __init__(self, html_text: str):
        super().__init__(convert_charrefs=False)
        self._raw = html_text
        self._line_starts = [0]
        for line in html_text.splitlines(keepends=True):
            self._line_starts.append(self._line_starts[-1] + len(line))
        self.slide = -1
        self.skipping = 0
        self.spans: list[dict] = []

    def _offset(self) -> int:
        line, col = self.getpos()
        return self._line_starts[line - 1] + col

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self.skipping += 1
        elif tag == "section" and any(name == "class" and "slide" in (value or "")
                                      for name, value in attrs):
            self.slide += 1

    def handle_endtag(self, tag):
        if tag in self.SKIP and self.skipping:
            self.skipping -= 1

    def handle_data(self, data):
        if self.skipping or self.slide < 0 or not data.strip():
            return
        start = self._offset()
        end = start + len(data)
        # 保守策略：原始切片与解析文本不一致（实体等）时跳过该节点
        if self._raw[start:end] != data:
            return
        self.spans.append({
            "slide": self.slide, "start": start, "end": end,
            "raw": data, "text": " ".join(data.split()),
        })


def iter_slide_texts(html_text: str) -> list[dict]:
    """Group visible text nodes per slide: [{index, title, texts:[{node, text}]}]."""
    mapper = _TextNodeMapper(html_text)
    mapper.feed(html_text)
    mapper.close()
    slides: dict[int, dict] = {}
    for position, span in enumerate(mapper.spans):
        slide = slides.setdefault(span["slide"], {
            "index": span["slide"], "title": "", "texts": []})
        slide["texts"].append({"slide": span["slide"], "node": position,
                               "text": span["text"]})
        if not slide["title"] and span["slide"] >= 0:
            slide["title"] = span["text"][:120]
    return [slides[key] for key in sorted(slides)]


def apply_text_edits(html_text: str, edits: list[dict]) -> str:
    """Replace exact text spans: edits = [{slide, node, text}]."""
    mapper = _TextNodeMapper(html_text)
    mapper.feed(html_text)
    mapper.close()
    result = html_text
    applied = 0
    for edit in sorted(edits, key=lambda item: -int(item.get("node", 0))):
        slide_index, node_index = int(edit.get("slide", -1)), int(edit.get("node", -1))
        new_text = str(edit.get("text") or "")
        matching = [span for span in mapper.spans
                    if span["slide"] == slide_index and mapper.spans.index(span) == node_index]
        if not matching:
            continue
        span = matching[0]
        start, end = span["start"], span["end"]
        if result[start:end] != span["raw"]:
            continue
        result = result[:start] + new_text + result[end:]
        applied += 1
    if not applied:
        raise RuntimeError("没有任何文本节点被更新（可能页面已变化），请刷新后重试")
    return result
